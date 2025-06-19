import os
from typing import List, Dict, Any
import PyPDF2
from docx import Document
import pandas as pd
import openpyxl
from pathlib import Path
import win32com.client
import pythoncom

class DocumentLoader:
    """文档加载器，支持多种文件格式的加载"""
    
    def __init__(self):
        self.supported_extensions = {
            '.pdf': self._load_pdf,
            '.docx': self._load_docx,
            '.doc': self._load_doc,
            '.xlsx': self._load_excel,
            '.xls': self._load_excel,
            '.txt': self._load_txt,
            '.csv': self._load_csv
        }
    
    def load_document(self, file_path: str) -> Dict[str, Any]:
        """
        加载单个文档
        
        Args:
            file_path: 文件路径
            
        Returns:
            包含文档内容的字典
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
            
        extension = file_path.suffix.lower()
        if extension not in self.supported_extensions:
            raise ValueError(f"不支持的文件格式: {extension}")
            
        loader = self.supported_extensions[extension]
        return loader(file_path)
    
    def load_directory(self, directory_path: str) -> List[Dict[str, Any]]:
        """
        加载目录中的所有支持格式的文档
        
        Args:
            directory_path: 目录路径
            
        Returns:
            文档内容列表
        """
        directory_path = Path(directory_path)
        if not directory_path.exists():
            raise FileNotFoundError(f"目录不存在: {directory_path}")
            
        documents = []
        for file_path in directory_path.rglob("*"):
            if file_path.is_file() and file_path.suffix.lower() in self.supported_extensions:
                try:
                    doc = self.load_document(str(file_path))
                    documents.append(doc)
                except Exception as e:
                    print(f"加载文件 {file_path} 时出错: {str(e)}")
                    
        return documents
    
    def _load_pdf(self, file_path: Path) -> Dict[str, Any]:
        """加载PDF文件"""
        content = []
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text = page.extract_text()
                    if text:
                        content.append(text)
                        
        except Exception as e:
            print(f"加载PDF文件失败: {str(e)}")
            raise
                        
        return {
            'file_path': str(file_path),
            'file_type': 'pdf',
            'content': '\n'.join(content)
        }
    
    def _load_docx(self, file_path: Path) -> Dict[str, Any]:
        """加载Word文档"""
        doc = Document(file_path)
        content = []
        for para in doc.paragraphs:
            if para.text.strip():
                content.append(para.text)
                
        return {
            'file_path': str(file_path),
            'file_type': 'docx',
            'content': '\n'.join(content)
        }
    
    def _load_excel(self, file_path: Path) -> Dict[str, Any]:
        """加载Excel文件"""
        content = []
        sheet_contents = {}  # 存储每个工作表的内容
        try:
            # 使用pandas读取所有工作表
            excel_file = pd.ExcelFile(file_path)
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                sheet_content = df.to_string()
                content.append(f"Sheet: {sheet_name}\n{sheet_content}")
                sheet_contents[sheet_name] = sheet_content
        except Exception as e:
            print(f"使用pandas加载Excel失败，尝试使用openpyxl: {str(e)}")
            # 如果pandas失败，使用openpyxl作为备选
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                sheet_content = []
                for row in ws.rows:
                    row_content = [str(cell.value) if cell.value is not None else '' for cell in row]
                    sheet_content.append('\t'.join(row_content))
                content_str = '\n'.join(sheet_content)
                content.append(f"Sheet: {sheet}\n" + content_str)
                sheet_contents[sheet] = content_str
                
        return {
            'file_path': str(file_path),
            'file_type': 'excel',
            'content': '\n\n'.join(content),
            'sheet_contents': sheet_contents  # 添加工作表内容字典
        }
    
    def get_sheet_content(self, doc_info: Dict[str, Any], sheet_name: str) -> str:
        """
        获取Excel文件中特定工作表的内容
        
        Args:
            doc_info: 文档信息字典
            sheet_name: 工作表名称
            
        Returns:
            工作表内容的字符串
        """
        if doc_info['file_type'] != 'excel':
            raise ValueError("只能从Excel文件中获取工作表内容")
            
        sheet_contents = doc_info.get('sheet_contents', {})
        if sheet_name not in sheet_contents:
            raise ValueError(f"工作表 '{sheet_name}' 不存在")
            
        return sheet_contents[sheet_name]
    
    def _load_txt(self, file_path: Path) -> Dict[str, Any]:
        """加载文本文件"""
        with open(file_path, 'r', encoding='utf-8') as file:
            content = file.read()
            
        return {
            'file_path': str(file_path),
            'file_type': 'txt',
            'content': content
        }
    
    def _load_csv(self, file_path: Path) -> Dict[str, Any]:
        """加载CSV文件"""
        try:
            df = pd.read_csv(file_path)
            content = df.to_string()
        except Exception as e:
            print(f"加载CSV文件失败: {str(e)}")
            content = ""
            
        return {
            'file_path': str(file_path),
            'file_type': 'csv',
            'content': content
        }
    
    def _load_doc(self, file_path: Path) -> Dict[str, Any]:
        """加载.doc格式的Word文档"""
        try:
            # 初始化COM
            pythoncom.CoInitialize()
            
            # 创建Word应用程序实例
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            
            try:
                # 打开文档
                doc = word.Documents.Open(str(file_path.absolute()))
                # 获取文本内容
                content = doc.Content.Text
                # 关闭文档
                doc.Close()
            finally:
                # 退出Word应用程序
                word.Quit()
                # 释放COM
                pythoncom.CoUninitialize()
            
            return {
                'file_path': str(file_path),
                'file_type': 'doc',
                'content': content
            }
        except Exception as e:
            print(f"加载.doc文件失败: {str(e)}")
            raise

# 使用示例
if __name__ == "__main__":
    loader = DocumentLoader()
    
    # 加载单个文件
    try:
        doc = loader.load_document("path/to/your/document.pdf")
        print(f"文件类型: {doc['file_type']}")
        print(f"文件内容: {doc['content'][:200]}...")  # 只打印前200个字符
    except Exception as e:
        print(f"加载文件时出错: {str(e)}")
    
    # 加载整个目录
    try:
        docs = loader.load_directory("path/to/your/documents")
        print(f"成功加载 {len(docs)} 个文档")
    except Exception as e:
        print(f"加载目录时出错: {str(e)}") 